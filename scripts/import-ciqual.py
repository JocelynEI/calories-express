"""Build the bundled subset from the official Ciqual 2025 workbook.
Usage: python3 scripts/import-ciqual.py /path/to/Ciqual2025.xlsx
Requires openpyxl for this one-off import, never in the app.
"""
import json
import re
import sys
from pathlib import Path
from openpyxl import load_workbook

def number(value):
    text = str(value).strip().replace(',', '.')
    # Unknown, traces and censored (<) measurements are not exact zeroes.
    return float(text) if re.fullmatch(r'\d+(?:\.\d+)?', text) else None

book = load_workbook(sys.argv[1], read_only=True, data_only=True)
rows = book['composition nutritionnelle'].iter_rows(values_only=True)
header = next(rows)
assert header[6] == 'alim_code' and header[7] == 'alim_nom_fr'
assert 'kcal' in header[10] and '6.25' in header[15].replace(',', '.')
foods = []
for row in rows:
    energy = number(row[10])
    if energy is None:
        continue
    foods.append([str(row[6]), row[7], energy, number(row[15]), number(row[16]), number(row[17])])
dest = Path(__file__).resolve().parents[1] / 'src/data/ciqual2025.json'
dest.write_text(json.dumps(foods, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
print(f'{len(foods)} foods → {dest.name}')
